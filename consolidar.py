import pandas as pd
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment

# 1. Encontrar todas as planilhas de vendas e juntar
arquivos = sorted(Path(".").glob("vendas_*.xlsx"))
lista = []
for arq in arquivos:
    df = pd.read_excel(arq)
    df["Mes"] = arq.stem.replace("vendas_", "").capitalize()
    lista.append(df)

todos = pd.concat(lista, ignore_index=True)

# 2. Calcular o total de cada venda
todos["Total"] = todos["Quantidade"] * todos["Preco"]

# 3. Resumo: total vendido por produto
resumo = (
    todos.groupby("Produto")["Total"]
    .sum()
    .reset_index()
    .sort_values("Total", ascending=False)
)

# 4. Salvar as duas abas no mesmo arquivo
with pd.ExcelWriter("relatorio.xlsx", engine="openpyxl") as writer:
    todos.to_excel(writer, sheet_name="Dados", index=False)
    resumo.to_excel(writer, sheet_name="Resumo", index=False)

# 5. Formatar com openpyxl
wb = load_workbook("relatorio.xlsx")

for ws in wb.worksheets:
    # Cabeçalho azul com letra branca e negrito
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(horizontal="center")

    # Formato de moeda nas colunas Preco e Total
    cabecalhos = [c.value for c in ws[1]]
    for idx, nome in enumerate(cabecalhos, start=1):
        if nome in ("Preco", "Total"):
            for linha in ws.iter_rows(min_row=2, min_col=idx, max_col=idx):
                for c in linha:
                    c.number_format = "R$ #,##0.00"

    # Ajustar largura das colunas ao conteúdo
    for col in ws.columns:
        largura = max(len(str(c.value)) if c.value is not None else 0 for c in col)
        ws.column_dimensions[col[0].column_letter].width = largura + 3

wb.save("relatorio.xlsx")
print("Relatório criado!")