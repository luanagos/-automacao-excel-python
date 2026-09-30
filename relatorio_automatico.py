import pandas as pd
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import CellIsRule
from datetime import datetime

# 1. Consolidar as planilhas de vendas
arquivos = sorted(Path(".").glob("vendas_*.xlsx"))
lista = [pd.read_excel(a).assign(Mes=a.stem.replace("vendas_", "").capitalize()) for a in arquivos]
todos = pd.concat(lista, ignore_index=True)
todos["Total"] = todos["Quantidade"] * todos["Preco"]

resumo = todos.groupby("Produto")["Total"].sum().reset_index().sort_values("Total", ascending=False)

# 2. Nome do arquivo com a data, para guardar um histórico
nome_arquivo = f"relatorio_{datetime.now().strftime('%Y-%m-%d')}.xlsx"

with pd.ExcelWriter(nome_arquivo, engine="openpyxl") as writer:
    todos.to_excel(writer, sheet_name="Dados", index=False)
    resumo.to_excel(writer, sheet_name="Resumo", index=False)

# 3. Formatação
wb = load_workbook(nome_arquivo)
for ws in wb.worksheets:
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(horizontal="center")
    for col in ws.columns:
        largura = max(len(str(c.value)) if c.value is not None else 0 for c in col)
        ws.column_dimensions[col[0].column_letter].width = largura + 3

# 4. Gráfico e formatação condicional na aba Resumo
ws = wb["Resumo"]
ultima = ws.max_row
verde = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
vermelho = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
ws.conditional_formatting.add(f"B2:B{ultima}", CellIsRule(operator="greaterThanOrEqual", formula=["10000"], fill=verde))
ws.conditional_formatting.add(f"B2:B{ultima}", CellIsRule(operator="lessThan", formula=["5000"], fill=vermelho))

grafico = BarChart()
grafico.title = "Total vendido por produto"
dados = Reference(ws, min_col=2, min_row=1, max_row=ultima)
categorias = Reference(ws, min_col=1, min_row=2, max_row=ultima)
grafico.add_data(dados, titles_from_data=True)
grafico.set_categories(categorias)
ws.add_chart(grafico, "D2")

wb.save(nome_arquivo)
print(f"Relatório {nome_arquivo} gerado com sucesso!")