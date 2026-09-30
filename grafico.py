from openpyxl import load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import PatternFill

wb = load_workbook("relatorio.xlsx")
ws = wb["Resumo"]
ultima = ws.max_row

# 1. Formatação condicional na coluna Total (B)
verde = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
vermelho = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")

ws.conditional_formatting.add(
    f"B2:B{ultima}", CellIsRule(operator="greaterThanOrEqual", formula=["10000"], fill=verde)
)
ws.conditional_formatting.add(
    f"B2:B{ultima}", CellIsRule(operator="lessThan", formula=["5000"], fill=vermelho)
)

# 2. Gráfico de barras
grafico = BarChart()
grafico.title = "Total vendido por produto"
grafico.y_axis.title = "Total (R$)"
grafico.x_axis.title = "Produto"

dados = Reference(ws, min_col=2, min_row=1, max_row=ultima)
categorias = Reference(ws, min_col=1, min_row=2, max_row=ultima)
grafico.add_data(dados, titles_from_data=True)
grafico.set_categories(categorias)
grafico.width = 16
grafico.height = 8

ws.add_chart(grafico, "D2")

wb.save("relatorio_final.xlsx")
print("Relatório final criado!")